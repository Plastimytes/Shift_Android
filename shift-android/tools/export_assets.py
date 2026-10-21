"""Regenerates app/src/main/assets/content/{categories,curriculum,sources}.json from docs/*.xlsx.
Run from repo root: python3 tools/export_assets.py   (needs openpyxl)"""
import json, re
from openpyxl import load_workbook
OUT = "app/src/main/assets/content/"
SHORT = {"A":"Japan","B":"Europe","C":"USA","D":"Racing","E":"Everyday","F":"Trucks","G":"Off-road","H":"Classic","I":"Powertrain","J":"Engineering"}
SEC = {"1 About Cars":"about-cars","2 Build the Car":"build-the-car","3 Repair and Modify":"repair-and-modify"}
slug = lambda s: re.sub(r"[^a-z0-9]+","-",s.lower()).strip("-")
bp = load_workbook("docs/SHIFT_Curriculum_Blueprint.xlsx", data_only=True)
groups, items = [], []
for g, name, cats, _ in bp["Interest Categories"].iter_rows(min_row=2, max_row=11, values_only=True):
    groups.append({"id": g, "name": name, "short": SHORT[g]})
    for c in [x.strip() for x in cats.split(";") if x.strip()]:
        items.append({"id": f"{g.lower()}-{slug(c)}", "name": c, "groupId": g})
assert len({i["id"] for i in items}) == len(items), "duplicate category ids"
json.dump({"groups": groups, "items": items}, open(OUT+"categories.json","w"), indent=1, ensure_ascii=False)
mods = []
for r in bp["Curriculum Blueprint"].iter_rows(min_row=2, values_only=True):
    if r[0] is None or r[1] not in SEC: break
    mods.append({"id": r[0], "section": SEC[r[1]], "title": r[2], "tier": r[3], "plannedLessons": r[7]})
json.dump({"modules": mods}, open(OUT+"curriculum.json","w"), indent=1, ensure_ascii=False)
reg = load_workbook("docs/SHIFT_Source_Registry.xlsx", data_only=True)["Source Registry"]
srcs = [{"id": r[0], "title": r[1], "issuer": r[2], "url": r[3], "tier": r[4], "reuse": r[8], "confidence": r[9]} for r in reg.iter_rows(min_row=2, values_only=True) if r[0]]
json.dump({"sources": srcs}, open(OUT+"sources.json","w"), indent=1, ensure_ascii=False)
print(len(groups), "groups,", len(items), "categories,", len(mods), "modules,", len(srcs), "sources")
