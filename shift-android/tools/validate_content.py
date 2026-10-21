"""SHIFT lesson checker.
Not a full JSON Schema engine (jsonschema could not be installed offline). It checks
the rules that matter most: required fields, id/enum patterns, strict-safety lessons
having a stop-and-refer block, every number carrying a source, region/review rules,
answer shapes, and that every cited source ID exists in the registry workbook.
Usage: python validate_content.py lesson.json [registry.xlsx]
"""
import json, re, sys
from openpyxl import load_workbook

def check(lesson, schema, registry_ids):
    errs = []
    for k in schema["required"]:
        if k not in lesson: errs.append(f"missing required field: {k}")
    if not re.match(schema["properties"]["id"]["pattern"], lesson.get("id","")): errs.append("bad lesson id")
    if not re.match(schema["properties"]["version"]["pattern"], lesson.get("version","")): errs.append("bad version")
    for f in ("section","difficulty"):
        if lesson.get(f) not in (schema["properties"][f].get("enum") or []): errs.append(f"bad {f}")
    for m in lesson.get("modes",[]):
        if m not in schema["properties"]["modes"]["items"]["enum"]: errs.append(f"bad mode {m}")
    blocks = lesson.get("blocks",[])
    btypes = [b.get("type") for b in blocks]
    saf = lesson.get("safety",{})
    if saf.get("level")=="strict" and "stop-and-refer" not in btypes: errs.append("strict safety lesson has no stop-and-refer block")
    if saf.get("level")=="strict" and saf.get("masteryThreshold",0) < 0.8: errs.append("strict lesson needs masteryThreshold >= 0.8")
    for b in blocks:
        if b.get("type")=="specification-table":
            for fg in b.get("figures",[]):
                for k in ("label","value","unit","sourceId","locator"):
                    if k not in fg: errs.append(f"figure missing {k}: {fg.get('label')}")
        if b.get("type")=="diagram" and len(b.get("altText",""))<10: errs.append("diagram altText too short")
    if lesson.get("jurisdiction",{}).get("scope")=="region-specific" and not lesson["jurisdiction"].get("regions"): errs.append("region-specific needs regions")
    v = lesson.get("vehicleSpecific",{})
    if v.get("isVehicleSpecific") and not v.get("specificationSourceId"): errs.append("vehicle-specific lesson needs specificationSourceId")
    r = lesson.get("review",{})
    if r.get("state") in ("expert-reviewed","published") and not (r.get("reviewer") and r.get("reviewedOn")): errs.append("reviewed state needs reviewer and date")
    for q in lesson.get("questions",[]):
        t=q.get("type")
        if t in ("multiple-choice","multi-select"):
            ids={o["id"] for o in q.get("options",[])}
            if not q.get("correctOptionIds") or not set(q["correctOptionIds"])<=ids: errs.append(f"{q.get('id')}: answer ids not in options")
        if t=="true-false" and "correctBoolean" not in q: errs.append(f"{q.get('id')}: missing correctBoolean")
        if t=="calculation" and "correctNumber" not in q: errs.append(f"{q.get('id')}: missing correctNumber")
        if not q.get("explanation") or not q.get("sourceIds"): errs.append(f"{q.get('id')}: needs explanation and sourceIds")
    cited=set()
    for b in blocks:
        for fg in b.get("figures",[]): cited.add(fg.get("sourceId"))
    for q in lesson.get("questions",[]): cited.update(q.get("sourceIds",[]))
    for m in lesson.get("misconceptions",[]): cited.update(m.get("sourceIds",[]))
    for ref in lesson.get("references",[]): cited.add(ref.get("sourceId"))
    if registry_ids is not None:
        for s in sorted(c for c in cited if c):
            if s not in registry_ids: errs.append(f"source {s} not in registry")
    refd={r["sourceId"] for r in lesson.get("references",[])}
    for s in cited:
        if s and s not in refd: errs.append(f"source {s} cited in body but missing from references list")
    return errs

if __name__=="__main__":
    lesson=json.load(open(sys.argv[1]))
    import os; schema=json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","content-src","shift_lesson.schema.json")))
    reg=None
    if len(sys.argv)>2:
        ws=load_workbook(sys.argv[2],read_only=True)["Source Registry"]
        reg={row[0] for row in ws.iter_rows(min_row=2,values_only=True) if row[0]}
    e=check(lesson,schema,reg)
    print("OK" if not e else "\n".join(e))
    sys.exit(1 if e else 0)
